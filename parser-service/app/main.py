import base64
import os
import secrets
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from fastapi import FastAPI, File, Header, HTTPException, UploadFile, status
from openpyxl import load_workbook
from pypdf import PdfReader


DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx"}

app = FastAPI(title="Sales AI Manager Parser Service")


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok", "service": "parser-service"}


@app.post("/parse")
async def parse_file(
    file: UploadFile = File(...),
    x_parser_token: str | None = Header(default=None),
) -> dict[str, Any]:
    verify_token(x_parser_token)

    filename = file.filename or "uploaded-file"
    extension = Path(filename).suffix.lower()
    if extension not in DOCUMENT_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Поддерживаются только .pdf, .docx и .xlsx.",
        )

    content = await file.read()
    if not content:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Файл пустой.")

    max_bytes = int(os.getenv("MAX_UPLOAD_MB", "15")) * 1024 * 1024
    if len(content) > max_bytes:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"Файл слишком большой. Максимальный размер: {os.getenv('MAX_UPLOAD_MB', '15')} МБ.",
        )

    try:
        if extension == ".pdf":
            parsed = parse_pdf(content)
        elif extension == ".docx":
            parsed = {"text": extract_docx(content), "pages": [], "parser": "python-docx"}
        else:
            parsed = {"text": extract_xlsx(content), "pages": [], "parser": "openpyxl"}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Не удалось прочитать файл. Проверьте формат и качество файла. Детали: {exc}",
        ) from exc

    parsed["filename"] = filename
    parsed["extension"] = extension
    parsed["text"] = (parsed.get("text") or "").strip()
    return parsed


def verify_token(value: str | None) -> None:
    expected = os.getenv("PARSER_SERVICE_TOKEN", "").strip()
    if not expected:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Parser service token не настроен.",
        )
    if value is None or not secrets.compare_digest(value, expected):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Неверный parser token.")


def parse_pdf(content: bytes) -> dict[str, Any]:
    pages: list[dict[str, Any]] = []
    warnings: list[str] = []

    try:
        import fitz  # PyMuPDF
    except Exception:
        text = extract_pdf_with_pypdf(content)
        return {
            "text": text,
            "pages": [],
            "parser": "pypdf",
            "warnings": ["PyMuPDF недоступен, изображения страниц PDF не созданы."],
        }

    max_image_pages = int(os.getenv("PDF_IMAGE_MAX_PAGES", "6"))
    image_dpi = int(os.getenv("PDF_IMAGE_DPI", "160"))

    doc = fitz.open(stream=content, filetype="pdf")
    try:
        extracted_text_parts: list[str] = []
        for index, page in enumerate(doc, start=1):
            page_text = (page.get_text("text") or "").strip()
            if page_text:
                extracted_text_parts.append(f"Страница {index}:\n{page_text}")
            pages.append({"page_number": index, "text": page_text})

        full_text = "\n\n".join(extracted_text_parts).strip()
        if not full_text:
            for index, page in enumerate(doc, start=1):
                if index > max_image_pages:
                    warnings.append(
                        f"PDF похож на скан. В vision-анализ переданы первые {max_image_pages} страниц."
                    )
                    break
                pix = page.get_pixmap(dpi=image_dpi)
                image_bytes = pix.tobytes("jpeg")
                encoded = base64.b64encode(image_bytes).decode("ascii")
                pages[index - 1].update(
                    {
                        "image_mime_type": "image/jpeg",
                        "image_base64": encoded,
                        "image_data_url": f"data:image/jpeg;base64,{encoded}",
                    }
                )

        return {"text": full_text, "pages": pages, "parser": "pymupdf", "warnings": warnings}
    finally:
        doc.close()


def extract_pdf_with_pypdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        page_text = (page.extract_text() or "").strip()
        if page_text:
            pages.append(f"Страница {index}:\n{page_text}")
    return "\n\n".join(pages)


def extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_rows: list[str] = []
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                table_rows.append(" | ".join(values))
    return "\n".join(paragraphs + table_rows)


def extract_xlsx(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    rows: list[str] = []
    for sheet in workbook.worksheets:
        rows.append(f"Лист: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                rows.append(" | ".join(values))
    return "\n".join(rows)
