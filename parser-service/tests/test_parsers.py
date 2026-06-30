from io import BytesIO

import pytest
from docx import Document
from fastapi import HTTPException
from openpyxl import Workbook

from app.main import extract_docx, extract_xlsx, parse_pdf, verify_token


def test_extract_docx_reads_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph("KBI Energy requests cable quote")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Qty 50"

    buffer = BytesIO()
    document.save(buffer)

    text = extract_docx(buffer.getvalue())

    assert "KBI Energy requests cable quote" in text
    assert "Item | Qty 50" in text


def test_extract_xlsx_reads_all_sheets() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Request"
    sheet.append(["Item", "Adapter RJ-45"])
    sheet.append(["Qty", 50])
    second = workbook.create_sheet("Delivery")
    second.append(["City", "Astana"])

    buffer = BytesIO()
    workbook.save(buffer)

    text = extract_xlsx(buffer.getvalue())

    assert "Лист: Request" in text
    assert "Item | Adapter RJ-45" in text
    assert "Qty | 50" in text
    assert "Лист: Delivery" in text
    assert "City | Astana" in text


def test_parse_pdf_reads_text_with_pymupdf() -> None:
    fitz = pytest.importorskip("fitz")
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Invoice 3818 for KBI Energy")
    content = document.tobytes()
    document.close()

    parsed = parse_pdf(content)

    assert parsed["parser"] == "pymupdf"
    assert "Invoice 3818 for KBI Energy" in parsed["text"]


def test_parse_blank_pdf_returns_image_page(monkeypatch: pytest.MonkeyPatch) -> None:
    fitz = pytest.importorskip("fitz")
    monkeypatch.setenv("PDF_IMAGE_MAX_PAGES", "1")
    document = fitz.open()
    document.new_page()
    content = document.tobytes()
    document.close()

    parsed = parse_pdf(content)

    assert parsed["text"] == ""
    assert parsed["pages"][0]["image_mime_type"] == "image/jpeg"
    assert parsed["pages"][0]["image_base64"]
    assert parsed["pages"][0]["image_data_url"].startswith("data:image/jpeg;base64,")


def test_verify_token_rejects_wrong_value(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("PARSER_SERVICE_TOKEN", "expected-token")

    with pytest.raises(HTTPException) as exc:
        verify_token("wrong-token")

    assert exc.value.status_code == 401
