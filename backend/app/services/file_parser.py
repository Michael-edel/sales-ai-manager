import base64
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import get_settings


DOCUMENT_EXTENSIONS = {".xlsx", ".pdf", ".docx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
AUDIO_EXTENSIONS = {".mp3", ".m4a", ".wav", ".ogg", ".opus", ".webm"}
SUPPORTED_EXTENSIONS = DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS | AUDIO_EXTENSIONS
IMAGE_MIME_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}


async def extract_text_from_upload(file: UploadFile) -> str:
    filename = file.filename or ""
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Поддерживаются файлы .xlsx, .pdf, .docx, .png, .jpg, .jpeg, .webp, .mp3, .m4a, .wav, .ogg, .opus и .webm.",
        )
    if extension in IMAGE_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Изображения нужно отправлять в vision-анализ, а не в текстовый парсер.",
        )

    content = await read_upload_bytes(file)

    text = extract_text_from_document_bytes(filename, content)

    text = text.strip()
    if not text:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Файл прочитан, но текст не найден. Возможно, PDF является сканом без OCR.",
        )
    return text


def extract_text_from_document_bytes(filename: str, content: bytes) -> str:
    extension = Path(filename or "").suffix.lower()
    if extension not in DOCUMENT_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Для извлечения текста поддерживаются только документы .xlsx, .pdf и .docx.",
        )

    try:
        if extension == ".xlsx":
            text = _extract_xlsx(content)
        elif extension == ".pdf":
            text = _extract_pdf(content)
        else:
            text = _extract_docx(content)
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Не удалось прочитать файл. Проверьте формат и качество файла. Детали: {exc}",
        ) from exc
    return text.strip()


def is_supported_image(filename: str) -> bool:
    return Path(filename or "").suffix.lower() in IMAGE_EXTENSIONS


def is_supported_audio(filename: str) -> bool:
    return Path(filename or "").suffix.lower() in AUDIO_EXTENSIONS


async def read_upload_bytes(file: UploadFile) -> bytes:
    content = await file.read()
    max_bytes = get_settings().max_upload_mb * 1024 * 1024
    if len(content) > max_bytes:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"Файл слишком большой. Максимальный размер: {get_settings().max_upload_mb} МБ.",
        )
    if not content:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Файл пустой.")
    return content


def image_to_data_url(filename: str, content: bytes) -> str:
    extension = Path(filename or "").suffix.lower()
    mime_type = IMAGE_MIME_TYPES.get(extension)
    if not mime_type:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Поддерживаются только изображения .png, .jpg, .jpeg и .webp.",
        )
    encoded = base64.b64encode(content).decode("ascii")
    return f"data:{mime_type};base64,{encoded}"


def _extract_xlsx(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    rows: list[str] = []
    for sheet in workbook.worksheets:
        rows.append(f"Лист: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                rows.append(" | ".join(values))
    return "\n".join(rows)


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(f"Страница {index}:\n{page_text}")
    return "\n\n".join(pages)


def _extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_rows = []
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                table_rows.append(" | ".join(values))
    return "\n".join(paragraphs + table_rows)
